"""Output-boundary regressions: truthful rows, stable IDs and read-back checks."""
from __future__ import annotations

import hashlib
import io
import json
import re
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from jobagent import output
from jobagent.drive_output import Publisher, DriveOutputError, SHEET_MIME
from jobagent.outreach import verification
from jobagent.profile import load_profile
from jobagent.runtime import now_iso
from tests.test_regressions import posting


@pytest.fixture
def profile():
    p = load_profile(Path(__file__).resolve().parents[1] / "profile.yaml.example")
    p.raw["identity"]["experience"]["years"] = 5
    p.raw["preferences"].update(home_country="India", onsite_cities=[["Hyderabad"]], qa_roles_only=True)
    return p


def good_row(**changes):
    text = "Build automated regression tests with Playwright and Python. " * 8
    row = {"Company": "Example", "Job Title": "QA Engineer", "Job Link": "https://jobs.ashbyhq.com/example/1",
           "JD Text": text, "JD Status": "verified_live", "JD Verified At": now_iso(),
           "JD SHA256": hashlib.sha256(text.encode()).hexdigest(), "Fit Status": "in_scope", "Rule Score": "80"}
    return {**row, **changes}


def report_summary():
    return {"checked_at": now_iso(), "status": "ok", "scope": "Test fixture", "checked": 1,
            "issues": [], "leads_not_checked": 0, "applications": "Not observed", "replies": "Not measured"}


def report_bytes(**changes):
    return output.make_workbook(output.partition([good_row(**changes)]), report_summary())


@pytest.mark.parametrize("change", [
    {"JD Text": "N/A"}, {"JD Verified At": "2000-01-01T00:00:00+00:00"},
    {"JD SHA256": "incorrect"}, {"Job Link": ""}, {"Job Title": ""}, {"JD Status": "blocked"},
])
def test_incomplete_or_stale_jobs_never_enter_actionable_tabs(change):
    tabs = output.partition([good_row(**change)])
    assert len(tabs["Failed checks"]) == 1
    assert not tabs["Ready to review"] and not tabs["Fit needs checking"]


def test_dedup_greenhouse_alias_and_application_url():
    rows = [{"Job Link": "https://boards.greenhouse.io/acme/jobs/12?source=search"},
            {"Job Link": "https://job-boards.greenhouse.io/acme/jobs/12"},
            {"Job Link": "https://jobs.ashbyhq.com/acme/123/application"},
            {"Job Link": "https://jobs.ashbyhq.com/acme/123"}]
    assert len(output.deduplicate(rows)) == 2


def test_excluded_location_not_promoted_by_experience_review(profile, monkeypatch):
    job = posting()
    job.location = "United States (US only)"
    job.description += " Must reside in the United States. Requires 9 years of QA experience."
    monkeypatch.setattr(verification, "fetch_board", lambda *args: ([job], "https://api.ashbyhq.com/example"))
    row = verification.verify_row({"Job Link": job.url}, profile, None)
    assert row["JD Status"] == "verified_live"
    assert row["Fit Status"] == "out_of_scope"


def test_experience_range_uses_lower_bound(profile, monkeypatch):
    job = posting()
    job.description += " Requires 3-5 years of QA experience."
    monkeypatch.setattr(verification, "fetch_board", lambda *args: ([job], "https://api.ashbyhq.com/example"))
    row = verification.verify_row({"Job Link": job.url}, profile, None)
    assert "3-5 years" in row["Experience Requirements (source excerpts)"]
    assert "5-year requirement" not in row["Fit Notes"]


def test_failed_board_is_fetched_once(profile, monkeypatch):
    calls = []
    def fail(*args):
        calls.append(1)
        raise ValueError("bad response")
    monkeypatch.setattr(verification, "fetch_board", fail)
    cache = {}
    for number in (1, 2):
        row = verification.verify_row({"Job Link": posting(number).url}, profile, None, board_cache=cache)
        assert row["JD Status"] == "fetch_error"
    assert len(calls) == 1


def test_medical_evaluator_is_not_a_software_qa_match(profile, monkeypatch):
    job = posting()
    job.title = "Dermatologist (AI Evaluation)"
    monkeypatch.setattr(verification, "fetch_board", lambda *args: ([job], "https://api.ashbyhq.com/example"))
    row = verification.verify_row({"Job Link": job.url}, profile, None)
    assert row["Fit Status"] == "out_of_scope"


def test_intermediary_does_not_enter_ready_list(profile, monkeypatch):
    job = posting()
    job.description += " Hiring on behalf of a partner company."
    monkeypatch.setattr(verification, "fetch_board", lambda *args: ([job], "https://api.ashbyhq.com/example"))
    row = verification.verify_row({"Job Link": job.url}, profile, None)
    assert row["Fit Status"] == "needs_review"
    assert row["Listing Type"].startswith("Intermediary")


def test_workbook_has_full_jd_and_treats_formula_like_text():
    value = '=HYPERLINK("https://untrusted.example", "text")'
    data = report_bytes(Company=value)
    wb = load_workbook(io.BytesIO(data), data_only=False)
    ws = wb["Ready to review"]
    assert ws["A2"].value == value and ws["A2"].data_type == "s"
    assert ws.cell(2, output.FIELDS.index("JD Text") + 1).value == good_row()["JD Text"]


def test_long_jd_survives_excel_cell_limit():
    text = "Actual source description. " * 3000
    data = report_bytes(**{"JD Text": text, "JD SHA256": hashlib.sha256(text.encode()).hexdigest()})
    ws = load_workbook(io.BytesIO(data))["Ready to review"]
    headers = [cell.value for cell in ws[1]]
    chunks = [ws.cell(2, i + 1).value for i, header in enumerate(headers) if header.startswith("JD Text")]
    assert "".join(chunks) == text
    assert all(len(chunk) <= 32000 for chunk in chunks)


def test_empty_search_and_database_produce_explicit_partial_report(tmp_path, profile):
    summary, rows = output.build_report(profile, [], tmp_path, limit=10)
    assert summary["status"] == "partial" and rows == []
    assert output.read_leads(tmp_path / "verified.csv") == []
    assert (tmp_path / "JobAgent Report.xlsx").exists()


class Request:
    def __init__(self, call):
        self.call = call
    def execute(self):
        return self.call()


class FakeDrive:
    def __init__(self):
        self.items = {}
        self.next_id = 1
        self.crash_after_create = False
        self.corrupt_export = False
        self.updates = 0
    def files(self):
        return self
    def list(self, q, **kwargs):
        role = re.search(r"key='role' and value='([^']+)'", q).group(1)
        return Request(lambda: {"files": [dict(v) for v in self.items.values() if v["appProperties"]["role"] == role]})
    def get(self, fileId, **kwargs):
        return Request(lambda: dict(self.items[fileId]))
    def create(self, body, media_body=None, **kwargs):
        def perform():
            key = str(self.next_id)
            self.next_id += 1
            data = media_body.getbytes(0, media_body.size()) if media_body else b""
            self.items[key] = {**body, "id": key, "data": data}
            if self.crash_after_create and body["mimeType"] == SHEET_MIME:
                self.crash_after_create = False
                raise TimeoutError("API committed but client lost response")
            return dict(self.items[key])
        return Request(perform)
    def update(self, fileId, media_body, **kwargs):
        def perform():
            self.updates += 1
            self.items[fileId]["data"] = media_body.getbytes(0, media_body.size())
            return {"id": fileId}
        return Request(perform)
    def export(self, fileId, **kwargs):
        return Request(lambda: report_bytes(Company="Corrupted") if self.corrupt_export else self.items[fileId]["data"])
    def get_media(self, fileId):
        return Request(lambda: self.items[fileId]["data"])


def bundle(tmp_path, data=None):
    (tmp_path / "JobAgent Report.xlsx").write_bytes(data or report_bytes())
    (tmp_path / "verification.json").write_text(json.dumps(report_summary()))
    (tmp_path / "verified.csv").write_text(output.csv_text([good_row()]))
    (tmp_path / "evidence.md").write_text("Test source evidence")


def test_republish_reuses_folder_and_sheet_ids(tmp_path):
    drive = FakeDrive()
    bundle(tmp_path)
    publisher = Publisher(drive, tmp_path / "state.json")
    first = publisher.publish(tmp_path)
    second = Publisher(drive, tmp_path / "state.json").publish(tmp_path)
    assert first["files"] == second["files"]
    assert len(drive.items) == 8
    bundle(tmp_path, report_bytes(Company="Updated"))
    third = Publisher(drive, tmp_path / "state.json").publish(tmp_path)
    assert first["sheet_id"] == third["sheet_id"] and len(drive.items) == 8


def test_uncertain_create_is_recovered_without_duplicate(tmp_path):
    drive = FakeDrive()
    drive.crash_after_create = True
    bundle(tmp_path)
    with pytest.raises(TimeoutError):
        Publisher(drive, tmp_path / "state.json").publish(tmp_path)
    result = Publisher(drive, tmp_path / "state.json").publish(tmp_path)
    assert result["status"] == "published_and_read_back" and len(drive.items) == 8


def test_application_tracker_is_never_overwritten(tmp_path):
    drive = FakeDrive()
    bundle(tmp_path)
    first = Publisher(drive, tmp_path / "state.json").publish(tmp_path)
    tracker = first["files"]["application_tracker"]
    drive.items[tracker]["data"] = b"User's actual application records"
    Publisher(drive, tmp_path / "state.json").publish(tmp_path)
    assert drive.items[tracker]["data"] == b"User's actual application records"


def test_user_edits_are_preserved(tmp_path):
    drive = FakeDrive()
    bundle(tmp_path)
    result = Publisher(drive, tmp_path / "state.json").publish(tmp_path)
    sheet = result["sheet_id"]
    drive.items[sheet]["data"] = report_bytes(Company="User's manual note")
    bundle(tmp_path, report_bytes(Company="New generated output"))
    with pytest.raises(DriveOutputError, match="unrecognized edits"):
        Publisher(drive, tmp_path / "state.json").publish(tmp_path)
    assert drive.updates == 0


def test_readback_mismatch_does_not_claim_publication(tmp_path):
    drive = FakeDrive()
    drive.corrupt_export = True
    bundle(tmp_path)
    with pytest.raises(DriveOutputError, match="read-back differs"):
        Publisher(drive, tmp_path / "state.json").publish(tmp_path)
    assert not (tmp_path / "publication.json").exists()


def test_pending_create_without_remote_evidence_does_not_retry(tmp_path):
    state = {"version": 1, "files": {"folder": {"create_pending": True}}}
    (tmp_path / "state.json").write_text(json.dumps(state))
    publisher = Publisher(FakeDrive(), tmp_path / "state.json")
    with pytest.raises(DriveOutputError, match="uncertain outcome"):
        publisher.ensure_folder()


def test_stale_bundle_rejected_before_drive_creation(tmp_path):
    bundle(tmp_path)
    (tmp_path / "verification.json").write_text(json.dumps({**report_summary(), "checked_at": "2000-01-01T00:00:00Z"}))
    drive = FakeDrive()
    with pytest.raises(DriveOutputError, match="24 hours"):
        Publisher(drive, tmp_path / "state.json").publish(tmp_path)
    assert drive.items == {}


def test_unattended_auth_failure_never_opens_browser(tmp_path, monkeypatch):
    from jobagent import drive_output
    from google_auth_oauthlib.flow import InstalledAppFlow
    token = tmp_path / "invalid-token.json"
    token.write_text("{}")
    monkeypatch.setattr(drive_output, "TOKEN_PATH", token)
    def unexpected(*args, **kwargs):
        pytest.fail("Unattended task must not open a consent browser")
    monkeypatch.setattr(InstalledAppFlow, "from_client_secrets_file", unexpected)
    with pytest.raises(DriveOutputError, match="reconnect"):
        drive_output.authenticate_drive("person@example.com")
