"""Send a daily status email summarizing a `morning` run's real results.

Reads logs/last-morning.json (written by run.py's main() for every `morning`
invocation) and only emails when real work actually completed this run -
a window/already-done no-op (status other than "completed") sends nothing,
so the two extra daily catch-up triggers that land outside the search
window don't spam an empty "success" email.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def load_json(path, default=None):
    path = Path(path)
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def ready_to_review_rows(xlsx_path, limit=15):
    from openpyxl import load_workbook

    path = Path(xlsx_path)
    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    if "Ready to review" not in wb.sheetnames:
        return []
    ws = wb["Ready to review"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = list(rows[0])
    idx = {name: i for i, name in enumerate(header)}
    out = []
    for row in rows[1:limit + 1]:
        out.append({
            "Company": row[idx["Company"]] if "Company" in idx else "",
            "Job Title": row[idx["Job Title"]] if "Job Title" in idx else "",
            "Job Link": row[idx["Job Link"]] if "Job Link" in idx else "",
        })
    return out


def build_body(details, rows):
    summary = details.get("summary") or {}
    publication = details.get("publication") or {}
    counts = summary.get("counts") or {}
    lines = [
        f"Job Agent - daily run - {details.get('date_ist', '')}",
        "",
        f"Vacancies checked: {summary.get('checked', 0)}",
        f"Ready to review: {counts.get('Ready to review', 0)}",
        f"Fit needs checking: {counts.get('Fit needs checking', 0)}",
        f"Excluded: {counts.get('Excluded jobs', 0)}",
        f"Startup shortlist: {summary.get('startup_shortlist_count', 0)}",
        "",
    ]
    drafts = summary.get("gmail_drafts") or {}
    if drafts:
        lines.append("Gmail drafts:")
        for status, count in drafts.items():
            lines.append(f"  {count}  {status}")
        lines.append("")
    if rows:
        lines.append("Ready to review:")
        for r in rows:
            lines.append(f"  {r['Company']} - {r['Job Title']}")
            lines.append(f"    {r['Job Link']}")
        lines.append("")
    sheet_url = publication.get("sheet_url")
    folder_url = publication.get("folder_url")
    if sheet_url:
        lines.append(f"Report: {sheet_url}")
    if folder_url:
        lines.append(f"Drive folder: {folder_url}")
    issues = summary.get("issues") or []
    if issues:
        lines.append("")
        lines.append(f"{len(issues)} source/research issue(s) this run - see the report for details.")
    return "\n".join(lines)


def main():
    details = load_json(ROOT / "logs" / "last-morning.json", {})
    if details.get("status") != "completed":
        print(f"No real work this run (status={details.get('status')!r}); no email sent.")
        return 0

    config = yaml.safe_load((ROOT / "config.yaml").read_text(encoding="utf-8"))
    expected_account = (config.get("drive_output") or {}).get("expected_account", "")
    if not expected_account:
        print("drive_output.expected_account not configured; skipping notification email.")
        return 0

    rows = ready_to_review_rows(ROOT / "research" / "report-latest" / "JobAgent Report.xlsx")
    body = build_body(details, rows)
    counts = (details.get("summary") or {}).get("counts") or {}
    subject = f"Job Agent - {details.get('date_ist', '')} - {counts.get('Ready to review', 0)} ready to review"

    from jobagent.notify import send_self_email
    result = send_self_email(subject, body, expected_account)
    print(f"Notification email sent: {result['message_id']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
