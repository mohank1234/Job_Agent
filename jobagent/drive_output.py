"""Publish app-owned reports through Drive, keeping permanent file IDs.

Google's Drive API can import/update XLSX as Google Sheets even when the Sheets
API is disabled. Only drive.file is requested in a separate OAuth token.
"""
from __future__ import annotations

import hashlib
import io
import json
import uuid
from pathlib import Path

from jobagent.runtime import atomic_json, now_iso, process_lock
from jobagent.output import XLSX_MIME, workbook_digest

DRIVE_SCOPE = "https://www.googleapis.com/auth/drive.file"
FOLDER_MIME = "application/vnd.google-apps.folder"
SHEET_MIME = "application/vnd.google-apps.spreadsheet"
ROOT = Path(__file__).resolve().parent.parent
TOKEN_PATH = ROOT / "drive_token.json"
STATE_PATH = ROOT / "drive_output_state.json"


class DriveOutputError(ValueError):
    pass


def authenticate_drive(expected_account, *, interactive=False):
    from google.auth.transport.requests import Request
    from google.auth.exceptions import RefreshError
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from googleapiclient.discovery import build

    if not expected_account:
        raise DriveOutputError("Configure drive_output.expected_account before publishing")
    creds = None
    try:
        if TOKEN_PATH.exists():
            creds = Credentials.from_authorized_user_file(str(TOKEN_PATH))
            if not creds.has_scopes([DRIVE_SCOPE]):
                creds = None
        if creds and not creds.valid and creds.refresh_token:
            creds.refresh(Request())
    except (ValueError, KeyError, RefreshError) as exc:
        if not interactive:
            raise DriveOutputError("Drive consent is invalid or expired. Run python run.py drive-auth to reconnect.") from exc
        creds = None
    if not creds or not creds.valid:
        if not interactive:
            raise DriveOutputError("Drive write consent is missing. Run python run.py drive-auth once; scheduled runs never open consent.")
        flow = InstalledAppFlow.from_client_secrets_file(str(ROOT / "credentials.json"), [DRIVE_SCOPE])
        creds = flow.run_local_server(
            port=0, timeout_seconds=600, login_hint=expected_account, prompt="consent",
            authorization_prompt_message="Complete Google Drive consent in the browser. If it did not open, use: {url}",
            success_message="Drive consent received. You can close this tab and return to Codex.",
        )
        granted = creds.granted_scopes or creds.scopes or []
        if DRIVE_SCOPE not in granted:
            raise DriveOutputError("Google Drive file access was not granted")
    drive = build("drive", "v3", credentials=creds, cache_discovery=False)
    account = drive.about().get(fields="user(emailAddress),importFormats").execute()
    if account["user"]["emailAddress"].casefold() != expected_account.casefold():
        raise DriveOutputError("Wrong Google account; no output will be created")
    if SHEET_MIME not in account.get("importFormats", {}).get(XLSX_MIME, []):
        raise DriveOutputError("This Drive account does not support XLSX-to-Sheets import")
    atomic_json(TOKEN_PATH, json.loads(creds.to_json()))
    return drive


def _escape(value):
    return value.replace("\\", "\\\\").replace("'", "\\'")


def deliverable_files(out):
    """An explicit output manifest prevents uploading authentication state/logs."""
    out = Path(out).resolve()
    manifest = out / "deliverables.json"
    if not manifest.exists():
        return []
    data = json.loads(manifest.read_text(encoding="utf-8"))
    if data.get("version") != 1 or not isinstance(data.get("files"), list) or len(data["files"]) > 50:
        raise DriveOutputError("Invalid deliverables manifest")
    allowed = {".csv": "text/csv", ".md": "text/markdown", ".json": "application/json",
               ".xlsx": XLSX_MIME, ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
               ".pdf": "application/pdf", ".txt": "text/plain"}
    reserved = {"JobAgent Report.xlsx", "verified.csv", "evidence.md", "verification.json", "publication.json", "deliverables.json"}
    seen, result = set(), []
    for name in data["files"]:
        if not isinstance(name, str) or Path(name).name != name or name in reserved or name in seen:
            raise DriveOutputError("Deliverables must have unique, unreserved filenames inside the output directory")
        path = out / name
        if path.is_symlink() or path.resolve().parent != out or not path.is_file() or path.suffix.lower() not in allowed:
            raise DriveOutputError(f"Invalid output deliverable: {name}")
        seen.add(name)
        result.append((path, "deliverable_" + hashlib.sha256(name.encode()).hexdigest()[:20], allowed[path.suffix.lower()]))
    result.append((manifest, "deliverables_manifest", "application/json"))
    return result


class Publisher:
    def __init__(self, drive, state_path=STATE_PATH):
        self.drive = drive
        self.state_path = Path(state_path)
        self.state = json.loads(self.state_path.read_text(encoding="utf-8")) if self.state_path.exists() else {"version": 1, "files": {}}
        if self.state.get("version") != 1 or not isinstance(self.state.get("files"), dict):
            raise DriveOutputError("Invalid output state; refusing to guess which Drive files to update")

    def save(self):
        atomic_json(self.state_path, self.state)

    def find(self, role, parent=None):
        query = ("trashed=false and appProperties has { key='jobagent_output' and value='v1' } "
                 f"and appProperties has {{ key='role' and value='{_escape(role)}' }}")
        if parent:
            query += f" and '{_escape(parent)}' in parents"
        files, token = [], None
        while True:
            result = self.drive.files().list(q=query, fields="nextPageToken,files(id,name,mimeType,parents,appProperties,trashed)",
                                             pageSize=100, pageToken=token).execute()
            files.extend(result.get("files", []))
            token = result.get("nextPageToken")
            if not token:
                break
        if len(files) > 1:
            raise DriveOutputError(f"Multiple managed {role} files found; resolve the duplicate IDs before publishing")
        return files[0] if files else None

    def existing(self, role, parent=None):
        entry = self.state["files"].setdefault(role, {})
        if entry.get("id"):
            meta = self.drive.files().get(fileId=entry["id"], fields="id,name,mimeType,parents,appProperties,trashed").execute()
        else:
            meta = self.find(role, parent)
        if meta:
            props = meta.get("appProperties", {})
            if meta.get("trashed") or props.get("jobagent_output") != "v1" or props.get("role") != role:
                raise DriveOutputError(f"The stored {role} ID is no longer a managed output file")
            if parent and parent not in meta.get("parents", []):
                raise DriveOutputError(f"The {role} file was moved out of the output folder; refusing to create another copy")
            entry["id"] = meta["id"]
            self.save()
        elif entry.get("create_pending"):
            raise DriveOutputError(f"Previous {role} creation has an uncertain outcome. Inspect Drive before clearing its pending state.")
        return meta

    def create(self, role, name, mime, *, parent=None, media=None):
        entry = self.state["files"].setdefault(role, {})
        entry.update(create_pending=True, creation_key=entry.get("creation_key") or uuid.uuid4().hex)
        self.save()
        body = {"name": name, "mimeType": mime, "appProperties": {
            "jobagent_output": "v1", "role": role, "creation_key": entry["creation_key"]}}
        if parent:
            body["parents"] = [parent]
        result = self.drive.files().create(body=body, media_body=media, fields="id,name,mimeType,parents,appProperties").execute()
        entry.update(id=result["id"], create_pending=False)
        self.save()
        return result

    def ensure_folder(self, name="JobAgent Output"):
        meta = self.existing("folder")
        if not meta:
            meta = self.create("folder", name, FOLDER_MIME)
        if meta["mimeType"] != FOLDER_MIME:
            raise DriveOutputError("Output folder ID does not identify a folder")
        return meta["id"]

    def ensure_tracker(self, folder_id):
        """Initialize once. The user owns all subsequent application entries."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
        from googleapiclient.http import MediaIoBaseUpload
        stored = self.state["files"].get("application_tracker", {})
        if stored.get("id"):
            prior = self.drive.files().get(fileId=stored["id"], fields="id,trashed,appProperties").execute()
            props = prior.get("appProperties", {})
            if (prior.get("trashed") and props.get("jobagent_output") == "v1"
                    and props.get("role") == "application_tracker"):
                # A removed optional tracker is a user choice. Do not restore it,
                # recreate a replacement, or block unrelated output publishing.
                return None
        meta = self.existing("application_tracker", folder_id)
        if meta:
            if meta["mimeType"] != SHEET_MIME:
                raise DriveOutputError("Application tracker ID is not a Sheet")
            return meta["id"]
        wb = Workbook()
        ws = wb.active
        ws.title = "Applications"
        ws.append(["Company", "Job Title", "Job Link", "Status", "Submitted At", "Confirmation Evidence",
                   "Reply At", "Reply Summary", "Next Action", "Notes"])
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(color="FFFFFF", bold=True)
            ws.column_dimensions[cell.column_letter].width = 32 if cell.column_letter not in ("C", "F", "J") else 60
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = "A1:J1000"
        validation = DataValidation(type="list", formula1='"Considering,Applied,Interview,Rejected,Offer,Withdrawn"', allow_blank=True)
        ws.add_data_validation(validation)
        validation.add("D2:D10000")
        help_sheet = wb.create_sheet("Instructions")
        help_sheet.append(["Application Tracker", "Your entries are preserved across report refreshes."])
        help_sheet.append(["Applied", "Record only after submission; include the confirmation reference, email or page evidence."])
        help_sheet.append(["Replies", "Record an actual employer response and date. An automatic acknowledgement is not an interview."])
        help_sheet.append(["Coverage", "This tracker starts empty. Existing emails and applications have not been imported."])
        help_sheet.column_dimensions["A"].width = 25
        help_sheet.column_dimensions["B"].width = 115
        stream = io.BytesIO()
        wb.save(stream)
        data = stream.getvalue()
        meta = self.create("application_tracker", "JobAgent - Application Tracker", SHEET_MIME, parent=folder_id,
                           media=MediaIoBaseUpload(io.BytesIO(data), mimetype=XLSX_MIME, resumable=False))
        remote = self.drive.files().export(fileId=meta["id"], mimeType=XLSX_MIME).execute()
        if workbook_digest(remote) != workbook_digest(data):
            raise DriveOutputError("Application tracker read-back mismatch")
        return meta["id"]

    def put_sheet(self, path, folder_id):
        from googleapiclient.http import MediaIoBaseUpload
        data = Path(path).read_bytes()
        digest = workbook_digest(data)
        role = "report"
        meta = self.existing(role, folder_id)
        entry = self.state["files"][role]
        if meta:
            if meta["mimeType"] != SHEET_MIME:
                raise DriveOutputError("Report ID is not a Google Sheet")
            remote = self.drive.files().export(fileId=meta["id"], mimeType=XLSX_MIME).execute()
            observed = workbook_digest(remote)
            if observed == digest:
                entry.update(grid_digest=digest, pending_digest=None, verified_at=now_iso())
                self.save()
                return meta["id"]
            allowed = {v for v in (entry.get("grid_digest"), entry.get("pending_digest")) if v}
            if observed not in allowed:
                raise DriveOutputError("The report contains unrecognized edits. Preserving them; review the report before replacing it.")
        entry["pending_digest"] = digest
        self.save()
        media = MediaIoBaseUpload(io.BytesIO(data), mimetype=XLSX_MIME, resumable=False)
        if meta:
            # Updating imported content preserves the file ID. Only this app's
            # generated report is managed; external/cloud reports are untouched.
            self.drive.files().update(fileId=meta["id"], body={"mimeType": SHEET_MIME}, media_body=media, fields="id").execute()
        else:
            meta = self.create(role, "JobAgent - Verified Jobs", SHEET_MIME, parent=folder_id, media=media)
        remote = self.drive.files().export(fileId=meta["id"], mimeType=XLSX_MIME).execute()
        if workbook_digest(remote) != digest:
            raise DriveOutputError("Drive upload returned, but report read-back differs; publication is not verified")
        entry.update(grid_digest=digest, pending_digest=None, verified_at=now_iso())
        self.save()
        return meta["id"]

    def put_file(self, path, role, mime, folder_id):
        from googleapiclient.http import MediaIoBaseUpload
        path = Path(path)
        data = path.read_bytes()
        digest = hashlib.sha256(data).hexdigest()
        meta = self.existing(role, folder_id)
        entry = self.state["files"][role]
        if meta:
            remote = self.drive.files().get_media(fileId=meta["id"]).execute()
            observed = hashlib.sha256(remote).hexdigest()
            if observed == digest:
                entry.update(content_digest=digest, pending_digest=None, verified_at=now_iso())
                self.save()
                return meta["id"]
            if observed not in {v for v in (entry.get("content_digest"), entry.get("pending_digest")) if v}:
                raise DriveOutputError(f"Preserving unrecognized edits in {path.name}; publication stopped")
        entry["pending_digest"] = digest
        self.save()
        media = MediaIoBaseUpload(io.BytesIO(data), mimetype=mime, resumable=False)
        if meta:
            self.drive.files().update(fileId=meta["id"], media_body=media, fields="id").execute()
        else:
            meta = self.create(role, path.name, mime, parent=folder_id, media=media)
        remote = self.drive.files().get_media(fileId=meta["id"]).execute()
        if hashlib.sha256(remote).hexdigest() != digest:
            raise DriveOutputError(f"Read-back mismatch for {path.name}")
        entry.update(content_digest=digest, pending_digest=None, verified_at=now_iso())
        self.save()
        return meta["id"]

    def publish(self, out, folder_name="JobAgent Output", *, main_files_only=False):
        out = Path(out)
        # Validate the complete local bundle before creating any remote files.
        summary = json.loads((out / "verification.json").read_text(encoding="utf-8"))
        from jobagent.outreach.service import recent
        if not recent(summary.get("checked_at"), max_age_days=1):
            raise DriveOutputError("Report is more than 24 hours old; refresh it before publishing")
        for name in ("JobAgent Report.xlsx", "verified.csv", "evidence.md"):
            if not (out / name).is_file():
                raise DriveOutputError(f"Report bundle is missing {name}")
        workbook_digest((out / "JobAgent Report.xlsx").read_bytes())
        extras = deliverable_files(out)
        if main_files_only:
            extras = [(p, r, m) for p, r, m in extras if 'resume' in p.name.casefold() and p.suffix.lower() in ('.docx', '.pdf')]
        folder = self.ensure_folder(folder_name)
        sheet = self.put_sheet(out / "JobAgent Report.xlsx", folder)
        files = {"report": sheet}
        for name, role, mime in ([] if main_files_only else [("verified.csv", "csv", "text/csv"),
                                 ("evidence.md", "evidence", "text/markdown"),
                                 ("verification.json", "status", "application/json")]):
            files[role] = self.put_file(out / name, role, mime, folder)
        files["report_xlsx"] = self.put_file(out / "JobAgent Report.xlsx", "report_xlsx", XLSX_MIME, folder)
        for path, role, mime in extras:
            files[role] = self.put_file(path, role, mime, folder)
        tracker = None if main_files_only else self.ensure_tracker(folder)
        if tracker:
            files["application_tracker"] = tracker
        result = {"status": "published_and_read_back", "verified_at": now_iso(),
                  "folder_id": folder, "sheet_id": sheet,
                  "folder_url": f"https://drive.google.com/drive/folders/{folder}",
                  "sheet_url": f"https://docs.google.com/spreadsheets/d/{sheet}/edit",
                  "tracker_url": f"https://docs.google.com/spreadsheets/d/{tracker}/edit" if tracker else None,
                  "tracker_status": "omitted in main-files mode" if main_files_only else "available" if tracker else "trashed; user deletion preserved", "files": files,
                  "main_files_only": main_files_only}
        atomic_json(out / "publication.json", result)
        # The receipt lists the verified deliverables; it does not recursively
        # include its own digest. Read-back still validates the receipt upload.
        if not main_files_only:
            self.put_file(out / "publication.json", "publication_receipt", "application/json", folder)
        self.state["last_publish"] = result
        self.save()
        return result


def publish_report(out, expected_account, folder_name="JobAgent Output", *, main_files_only=False):
    drive = authenticate_drive(expected_account)
    with process_lock(STATE_PATH.with_suffix(".lock")):
        return Publisher(drive).publish(out, folder_name, main_files_only=main_files_only)
